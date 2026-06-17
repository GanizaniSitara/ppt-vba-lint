"""
drawio_ppt.py — Bidirectional converter between Draw.io XML and PowerPoint.

Usage:
    python drawio_ppt.py drawio2ppt  input.drawio  output.pptx  [--config config.ini]
    python drawio_ppt.py drawio2ppt-fidelity input.drawio output.pptx
    python drawio_ppt.py drawio2odp  input.drawio  output.odp   [--config config.ini]
    python drawio_ppt.py ppt2drawio  input.pptx    output.drawio

drawio2ppt: Parses .drawio XML (or raw mxGraphModel XML), creates a .pptx
            with native shapes, connectors, colors, and labels.  Optionally
            snaps colours to a corporate palette via --config.

drawio2ppt-fidelity: Preserves visible Draw.io geometry as editable
                     PowerPoint shapes/text/connectors, including hidden-layer
                     filtering, opacity blending, dense labels, and arrowheads.

drawio2odp: Parses .drawio XML and writes an OpenDocument Presentation (.odp)
            package directly, without PowerPoint, LibreOffice, or COM.

ppt2drawio: Reads a .pptx, extracts every shape/connector on every slide,
            and writes a .drawio file (one diagram page per slide).
"""

from __future__ import annotations

import argparse
import base64
import configparser
from dataclasses import dataclass
from datetime import datetime, timezone
import math
import re
import sys
import zipfile
import zlib
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional
from urllib.parse import unquote
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR_TYPE, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt


# ---------------------------------------------------------------------------
# Drawio encoding helpers (ported from drawio_tools.py)
# ---------------------------------------------------------------------------

def decode_diagram_data(data: str) -> str:
    raw = base64.b64decode(data)
    raw = zlib.decompressobj(-15).decompress(raw)
    return unquote(raw.decode())


def encode_diagram_data(xml: str) -> str:
    from urllib.parse import quote
    encoded = quote(xml, safe="~()*!.'")
    compressed = zlib.compressobj(
        zlib.Z_DEFAULT_COMPRESSION, zlib.DEFLATED, -15,
        memLevel=8, strategy=zlib.Z_DEFAULT_STRATEGY,
    ).compress(encoded.encode())
    compressed += zlib.compressobj(
        zlib.Z_DEFAULT_COMPRESSION, zlib.DEFLATED, -15,
        memLevel=8, strategy=zlib.Z_DEFAULT_STRATEGY,
    ).flush()
    # simpler: do it in one shot
    c = zlib.compressobj(zlib.Z_DEFAULT_COMPRESSION, zlib.DEFLATED, -15)
    out = c.compress(encoded.encode()) + c.flush()
    return base64.b64encode(out).decode()


# ---------------------------------------------------------------------------
# HTML-to-text helper
# ---------------------------------------------------------------------------

class _StripHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self._parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in ("br", "div", "p"):
            self._parts.append("\n")

    def handle_data(self, data):
        self._parts.append(data)

    def get_text(self) -> str:
        # Collapse multiple newlines, strip leading/trailing
        import re
        text = "".join(self._parts)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


def strip_html(html: str) -> str:
    p = _StripHTML()
    p.feed(html)
    return p.get_text()


# ---------------------------------------------------------------------------
# Style parsing
# ---------------------------------------------------------------------------

def parse_style(style_str: str) -> dict[str, str]:
    """Parse a drawio style string like 'rounded=1;fillColor=#dae8fc;...' into a dict."""
    result: dict[str, str] = {}
    if not style_str:
        return result
    for token in style_str.rstrip(";").split(";"):
        if "=" in token:
            k, v = token.split("=", 1)
            result[k.strip()] = v.strip()
        elif token.strip():
            result[token.strip()] = ""
    return result


def build_style(props: dict[str, str]) -> str:
    parts = []
    for k, v in props.items():
        parts.append(f"{k}={v}" if v else k)
    return ";".join(parts) + ";"


# ---------------------------------------------------------------------------
# Colour helpers
# ---------------------------------------------------------------------------

def hex_to_rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    if len(h) < 6:
        return (0, 0, 0)
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def rgb_to_hex(r: int, g: int, b: int) -> str:
    return f"#{r:02X}{g:02X}{b:02X}"


def colour_distance(c1: tuple[int, int, int], c2: tuple[int, int, int]) -> float:
    return math.sqrt(sum((a - b) ** 2 for a, b in zip(c1, c2)))


def load_palette(ini_path: str) -> list[str]:
    """Load colour hex codes from config.ini [colours] section."""
    cfg = configparser.ConfigParser()
    cfg.read(ini_path)
    colours = []
    if cfg.has_section("colours"):
        for key in cfg["colours"]:
            val = cfg["colours"][key]
            hex_code = val.split("|")[0].strip()
            colours.append(hex_code)
    return colours


def snap_colour(hex_colour: str, palette: list[str]) -> str:
    if not palette:
        return hex_colour
    target = hex_to_rgb(hex_colour)
    best = min(palette, key=lambda c: colour_distance(target, hex_to_rgb(c)))
    return best


def darken(hex_colour: str, amount: int = 20) -> str:
    r, g, b = hex_to_rgb(hex_colour)
    return rgb_to_hex(max(0, r - amount), max(0, g - amount), max(0, b - amount))


def is_hex_colour(value: str) -> bool:
    value = value.strip().lstrip("#")
    return bool(re.fullmatch(r"[0-9A-Fa-f]{6}", value))


# ---------------------------------------------------------------------------
# Drawio XML parsing
# ---------------------------------------------------------------------------

def load_drawio_xml(path: str) -> ET.Element:
    """Load a .drawio or raw mxGraphModel XML file. Returns the mxGraphModel element."""
    tree = ET.parse(path)
    root = tree.getroot()

    # If it's a full .drawio (mxfile wrapper), decode diagram data
    if root.tag == "mxfile":
        diagram = root.find("diagram")
        if diagram is not None:
            text = (diagram.text or "").strip()
            if text:
                # Compressed/encoded diagram
                try:
                    xml_str = decode_diagram_data(text)
                    return ET.fromstring(xml_str)
                except Exception:
                    pass
            # Uncompressed — mxGraphModel is a child element
            mg = diagram.find("mxGraphModel")
            if mg is not None:
                return mg

    # Raw mxGraphModel XML (like taxonomy.drawio.xml)
    if root.tag == "mxGraphModel":
        return root

    raise ValueError(f"Cannot find mxGraphModel in {path}")


def get_absolute_coords(nodes_by_id: dict, node_el: ET.Element) -> tuple[float, float]:
    """Walk the parent chain to accumulate absolute x, y."""
    x, y = 0.0, 0.0
    current = node_el
    while current is not None:
        geo = current.find("mxGeometry")
        if geo is not None:
            x += float(geo.get("x", 0))
            y += float(geo.get("y", 0))
        pid = current.get("parent", "0")
        if pid == "0" or pid not in nodes_by_id:
            break
        current = nodes_by_id[pid]
    return x, y


def resolve_label(node_el: ET.Element, parent_el: Optional[ET.Element]) -> str:
    """Extract the human-readable label, resolving %placeholder% tokens."""
    parts = []

    # Label from <object> parent
    if parent_el is not None and parent_el.tag == "object":
        lbl = parent_el.get("label", "")
        if lbl:
            parts.append(strip_html(lbl))

    # Value from the mxCell itself
    val = node_el.get("value", "")
    if val:
        parts.append(strip_html(val))

    label = "\n".join(p for p in parts if p)

    # Resolve %placeholder% tokens from the <object> parent
    if "%" in label and parent_el is not None:
        def _replace(m):
            key = m.group(1)
            return parent_el.get(key, key)
        label = re.sub(r"%(\w+)%", _replace, label)

    # Append description if present
    desc_el = parent_el if parent_el is not None else node_el
    desc = desc_el.get("description", "")
    if desc:
        label = label + "\n" + desc.strip()

    return label


# ---------------------------------------------------------------------------
# drawio → pptx
# ---------------------------------------------------------------------------

# Drawio uses pixels; PowerPoint uses EMU (914400 per inch, 12700 per pt).
# Drawio's coordinate system is roughly 1 pixel = 1 CSS pixel ≈ 0.75pt.
DRAWIO_PX_TO_EMU = int(0.75 * 12700)  # 9525 EMU per drawio pixel


def dx(val: float) -> int:
    return int(val * DRAWIO_PX_TO_EMU)


def drawio_to_pptx(
    drawio_path: str,
    pptx_path: str,
    config_path: Optional[str] = None,
    exclude_layers: Optional[list[str]] = None,
    back_layers: Optional[list[str]] = None,
    line_darken: int = 20,
) -> None:
    palette = load_palette(config_path) if config_path else []
    exclude_layers = [l.lower() for l in (exclude_layers or [])]
    back_layers = [l.lower() for l in (back_layers or [])]

    mg = load_drawio_xml(drawio_path)
    root_el = mg.find("root")
    if root_el is None:
        raise ValueError("No <root> element in mxGraphModel")

    # Index all elements by id
    all_elements: dict[str, ET.Element] = {}
    for el in root_el.iter():
        eid = el.get("id")
        if eid:
            all_elements[eid] = el

    # Build parent map (child → parent element)
    parent_map: dict[str, ET.Element] = {}
    for el in root_el.iter():
        for child in el:
            cid = child.get("id")
            if cid:
                parent_map[cid] = el

    # Determine which IDs belong to excluded layers
    def is_excluded(el: ET.Element) -> bool:
        current = el
        while current is not None:
            val = (current.get("value") or "").strip().lower()
            if val and val in exclude_layers:
                return True
            pid = current.get("parent", "0")
            if pid == "0" or pid not in all_elements:
                break
            current = all_elements[pid]
        return False

    def is_back_layer(el: ET.Element) -> bool:
        current = el
        while current is not None:
            val = (current.get("value") or "").strip().lower()
            if val and val in back_layers:
                return True
            pid = current.get("parent", "0")
            if pid == "0" or pid not in all_elements:
                break
            current = all_elements[pid]
        return False

    # Collect vertices and edges, tracking which <object> wraps each mxCell.
    # mxCells inside <object> elements inherit the object's id and attributes.
    vertices: list[tuple[ET.Element, Optional[ET.Element]]] = []  # (mxCell, object_parent)
    edges: list[tuple[ET.Element, Optional[ET.Element]]] = []
    seen_cells: set[int] = set()  # avoid duplicates by element identity

    # First pass: mxCells inside <object> wrappers
    for obj in root_el.iter("object"):
        if is_excluded(obj):
            continue
        for cell in obj.findall("mxCell"):
            if id(cell) in seen_cells:
                continue
            seen_cells.add(id(cell))
            if cell.get("vertex") == "1":
                vertices.append((cell, obj))
            elif cell.get("edge") == "1":
                edges.append((cell, obj))

    # Second pass: standalone mxCells (not inside objects)
    for el in root_el.iter("mxCell"):
        if id(el) in seen_cells:
            continue
        seen_cells.add(id(el))
        if is_excluded(el):
            continue
        if el.get("vertex") == "1":
            vertices.append((el, None))
        elif el.get("edge") == "1":
            edges.append((el, None))

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 16:9
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    shape_map: dict[str, object] = {}  # drawio id → pptx shape

    # --- bounding box pass ---
    bb_min_x, bb_min_y = float("inf"), float("inf")
    bb_max_x, bb_max_y = float("-inf"), float("-inf")
    for v, _obj in vertices:
        geo = v.find("mxGeometry")
        if geo is None:
            continue
        ax, ay = get_absolute_coords(all_elements, v)
        w = float(geo.get("width", 50))
        h = float(geo.get("height", 50))
        bb_min_x = min(bb_min_x, ax)
        bb_min_y = min(bb_min_y, ay)
        bb_max_x = max(bb_max_x, ax + w)
        bb_max_y = max(bb_max_y, ay + h)

    # NOTE: Don't include edge mxPoints in bounding box — they can be wildly
    # off-screen and distort the layout. Vertices define the real diagram extent.

    diag_w = bb_max_x - bb_min_x or 1
    diag_h = bb_max_y - bb_min_y or 1
    margin_emu = Pt(30)  # 30pt margin (~10mm)
    avail_w = prs.slide_width - 2 * margin_emu
    avail_h = prs.slide_height - 2 * margin_emu
    scale_x = avail_w / dx(diag_w)
    scale_y = avail_h / dx(diag_h)
    scale = min(scale_x, scale_y, 1.0)  # don't upscale beyond 1:1

    # Center the diagram on the slide
    rendered_w = int(dx(diag_w) * scale)
    rendered_h = int(dx(diag_h) * scale)
    offset_x = (prs.slide_width - rendered_w) // 2
    offset_y = (prs.slide_height - rendered_h) // 2

    def to_emu_x(val: float) -> int:
        return int((dx(val - bb_min_x)) * scale + offset_x)

    def to_emu_y(val: float) -> int:
        return int((dx(val - bb_min_y)) * scale + offset_y)

    def to_emu_w(val: float) -> int:
        return int(dx(val) * scale)

    # --- Sort vertices: largest shapes first (containers behind), smallest on top ---
    def _shape_area(item):
        v, _ = item
        geo = v.find("mxGeometry")
        if geo is None:
            return 0
        return float(geo.get("width", 50)) * float(geo.get("height", 50))

    vertices.sort(key=lambda item: (0 if is_back_layer(item[0]) else 1, -_shape_area(item)))

    # --- Classify shapes ---
    MIN_SHAPE_AREA = 1500  # skip tiny annotation shapes (< ~39x39 drawio px)
    from pptx.enum.text import PP_ALIGN
    import re as _re

    rendered_count = 0
    skipped_count = 0

    for v, obj_parent in vertices:
        geo = v.find("mxGeometry")
        if geo is None:
            continue

        ax, ay = get_absolute_coords(all_elements, v)
        w = float(geo.get("width", 50))
        h = float(geo.get("height", 50))
        style = parse_style(v.get("style", ""))
        label = resolve_label(v, obj_parent)

        # Skip ellipses entirely — at slide scale they're 60px circles crammed together
        # and create an unreadable cluster. Also skip empty-label shapes < threshold.
        if "ellipse" in style:
            skipped_count += 1
            continue
        if w * h < MIN_SHAPE_AREA and not label.strip():
            skipped_count += 1
            continue

        vid = (obj_parent.get("id", "") if obj_parent is not None else "") or v.get("id", "")

        # Detect text-only elements (annotations, notes)
        is_text_only = "text" in style and style.get("fillColor", "none").lower() in ("none", "")

        # Shape type
        if "ellipse" in style:
            shape_type = MSO_SHAPE.OVAL
        elif style.get("rounded") == "1":
            shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            shape_type = MSO_SHAPE.RECTANGLE

        # Create shape
        sx, sy, sw, sh = to_emu_x(ax), to_emu_y(ay), to_emu_w(w), to_emu_w(h)
        if is_text_only:
            shp = slide.shapes.add_textbox(sx, sy, sw, sh)
        else:
            shp = slide.shapes.add_shape(shape_type, sx, sy, sw, sh)

        # --- Text ---
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(3)
        tf.margin_right = Pt(3)
        tf.margin_top = Pt(2)
        tf.margin_bottom = Pt(2)
        tf.auto_size = None  # let PowerPoint handle overflow
        tf.text = label

        # Font size — 8pt base, let auto-shrink handle overflow
        font_size = 8
        if "ellipse" in style:
            font_size = 6

        # Font colour from drawio style
        font_color_str = style.get("fontColor", "#000000")
        fr, fg, fb = 0, 0, 0
        if font_color_str and font_color_str.startswith("#"):
            fr, fg, fb = hex_to_rgb(font_color_str)
        elif font_color_str and "rgb" in font_color_str:
            nums = _re.findall(r"\d+", font_color_str)
            if len(nums) >= 3:
                fr, fg, fb = int(nums[0]), int(nums[1]), int(nums[2])

        # Parse drawio font sizes from the HTML label
        # Drawio C4 labels have: <font style="font-size: 16px"><b>Name</b></font>
        # followed by [Type] and <font style="font-size: 11px" color="#cccccc">Description</font>
        raw_html = ""
        if obj_parent is not None and obj_parent.tag == "object":
            raw_html = obj_parent.get("label", "")
        if not raw_html:
            raw_html = v.get("value", "") or ""

        # Extract font-size hints from the HTML for different text segments
        import re as _re3
        font_sizes_in_html = _re3.findall(r'font-size:\s*(\d+)px', raw_html)

        lines = label.split("\n")
        non_empty = [l.strip() for l in lines if l.strip()]
        tf.clear()

        align_val = style.get("align", "center")

        for i, line in enumerate(lines):
            line = line.strip()
            if not line:
                continue
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            # Center-align by default (C4 style)
            if align_val == "left":
                p.alignment = PP_ALIGN.LEFT
            else:
                p.alignment = PP_ALIGN.CENTER

            run = p.add_run()
            run.text = line

            # Size the text based on what segment this is
            # First non-empty line = name (large, bold)
            # Lines starting with [ = type (medium)
            # Other lines = description (smaller, lighter)
            line_idx_in_content = non_empty.index(line) if line in non_empty else i

            if line_idx_in_content == 0:
                # Name line — use drawio's font-size if available, else scale
                name_size = float(font_sizes_in_html[0]) if font_sizes_in_html else 16
                run.font.size = Pt(name_size * scale * 0.75)
                run.font.bold = True
                run.font.color.rgb = RGBColor(fr, fg, fb)
            elif line.startswith("["):
                # Type line — regular size
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(fr, fg, fb)
            else:
                # Description — smaller, lighter
                desc_size = float(font_sizes_in_html[-1]) if len(font_sizes_in_html) > 1 else 11
                run.font.size = Pt(desc_size * scale * 0.75)
                # If base font is light (white), make description dimmer
                if fr > 200 and fg > 200 and fb > 200:
                    run.font.color.rgb = RGBColor(204, 204, 204)
                else:
                    run.font.color.rgb = RGBColor(min(fr + 60, 255), min(fg + 60, 255), min(fb + 60, 255))

        # Vertical text alignment — default to middle unless drawio says top
        va = style.get("verticalAlign", "middle")
        if not is_text_only and hasattr(shp, "text_frame"):
            from pptx.enum.text import MSO_ANCHOR
            if va == "top":
                shp.text_frame.paragraphs[0]  # ensure exists
                # Set anchor via the bodyPr element
                try:
                    from pptx.oxml.ns import qn as _qn2
                    bodyPr = shp.text_frame._txBody.find(_qn2("a:bodyPr"))
                    if bodyPr is not None:
                        bodyPr.set("anchor", "t")
                except Exception:
                    pass
            else:
                try:
                    from pptx.oxml.ns import qn as _qn2
                    bodyPr = shp.text_frame._txBody.find(_qn2("a:bodyPr"))
                    if bodyPr is not None:
                        bodyPr.set("anchor", "ctr")
                except Exception:
                    pass

        if is_text_only:
            shape_map[vid] = shp
            rendered_count += 1
            continue

        # --- Fill ---
        fill_hex = style.get("fillColor", "")
        if fill_hex and fill_hex.lower() != "none":
            if palette:
                fill_hex = snap_colour(fill_hex, palette)
            r, g, b = hex_to_rgb(fill_hex)
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor(r, g, b)
        elif fill_hex.lower() == "none":
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor(255, 255, 255)

        # --- Stroke ---
        stroke_hex = style.get("strokeColor", "")
        if stroke_hex and stroke_hex.startswith("#"):
            if line_darken > 0 and palette:
                stroke_hex = darken(stroke_hex, line_darken)
            r, g, b = hex_to_rgb(stroke_hex)
            shp.line.color.rgb = RGBColor(r, g, b)
        shp.line.width = Pt(0.75)

        if style.get("dashed") == "1":
            shp.line.dash_style = 4

        shape_map[vid] = shp
        rendered_count += 1

    # --- Connectors ---
    # For simple diagrams (few shapes), allow shape-to-shape connectors via edge intersection.
    # For complex diagrams (many shapes), only render edges with explicit mxPoints to avoid noise.
    simple_diagram = rendered_count <= 30
    connector_count = 0
    from pptx.enum.shapes import MSO_CONNECTOR_TYPE

    def _edge_intersect(shp, other_cx, other_cy):
        """Line from shape center to target, clipped at shape rectangle boundary."""
        cx = shp.left + shp.width // 2
        cy = shp.top + shp.height // 2
        ddx = other_cx - cx
        ddy = other_cy - cy
        if ddx == 0 and ddy == 0:
            return cx, cy
        hw = shp.width // 2
        hh = shp.height // 2
        s_x = abs(hw / ddx) if ddx != 0 else float("inf")
        s_y = abs(hh / ddy) if ddy != 0 else float("inf")
        s = min(s_x, s_y)
        return int(cx + ddx * s), int(cy + ddy * s)

    for e, _edge_obj in edges:
        geo = e.find("mxGeometry")
        if geo is None:
            continue

        source_id = e.get("source", "")
        target_id = e.get("target", "")
        src_pt = geo.find("mxPoint[@as='sourcePoint']")
        tgt_pt = geo.find("mxPoint[@as='targetPoint']")

        def _best_edge_pair(s_shp, t_shp):
            """Pick the best source/target edge midpoints for a connector.
            Uses the target's center-x for the source exit point when going vertical,
            so the line goes straight down to the target."""
            sc = (s_shp.left + s_shp.width // 2, s_shp.top + s_shp.height // 2)
            tc = (t_shp.left + t_shp.width // 2, t_shp.top + t_shp.height // 2)
            dx_abs = abs(tc[0] - sc[0])
            dy_abs = abs(tc[1] - sc[1])
            if dy_abs >= dx_abs * 0.5:
                # Primarily vertical — exit from source bottom center, enter target top center
                if sc[1] < tc[1]:
                    # Source exit: use target's x so the line is more vertical
                    src_x = min(max(tc[0], s_shp.left), s_shp.left + s_shp.width)
                    return (src_x, s_shp.top + s_shp.height), (tc[0], t_shp.top)
                else:
                    src_x = min(max(tc[0], s_shp.left), s_shp.left + s_shp.width)
                    return (src_x, s_shp.top), (tc[0], t_shp.top + t_shp.height)
            else:
                # Primarily horizontal
                if sc[0] < tc[0]:
                    return (s_shp.left + s_shp.width, sc[1]), (t_shp.left, tc[1])
                else:
                    return (s_shp.left, sc[1]), (t_shp.left + t_shp.width, tc[1])

        # Resolve endpoints
        if src_pt is not None:
            sx_emu = to_emu_x(float(src_pt.get("x", 0)))
            sy_emu = to_emu_y(float(src_pt.get("y", 0)))
            if tgt_pt is not None:
                tx_emu = to_emu_x(float(tgt_pt.get("x", 0)))
                ty_emu = to_emu_y(float(tgt_pt.get("y", 0)))
            elif simple_diagram and target_id and target_id in shape_map:
                t = shape_map[target_id]
                tx_emu, ty_emu = _edge_intersect(t, sx_emu, sy_emu)
            else:
                continue
        elif simple_diagram and source_id in shape_map:
            s = shape_map[source_id]
            if target_id and target_id in shape_map:
                t = shape_map[target_id]
                (sx_emu, sy_emu), (tx_emu, ty_emu) = _best_edge_pair(s, t)
            elif tgt_pt is not None:
                tx_emu = to_emu_x(float(tgt_pt.get("x", 0)))
                ty_emu = to_emu_y(float(tgt_pt.get("y", 0)))
                sx_emu, sy_emu = _edge_intersect(s, tx_emu, ty_emu)
            else:
                continue
        else:
            continue

        edge_style = parse_style(e.get("style", ""))

        # Use elbow (right-angle) connector if drawio specifies orthogonal routing
        is_orthogonal = "orthogonal" in edge_style.get("edgeStyle", "")
        conn_type = MSO_CONNECTOR_TYPE.ELBOW if is_orthogonal else MSO_CONNECTOR_TYPE.STRAIGHT
        conn = slide.shapes.add_connector(
            conn_type,
            sx_emu, sy_emu, tx_emu, ty_emu,
        )
        stroke = edge_style.get("strokeColor", "")
        if stroke and stroke.startswith("#"):
            er, eg, eb = hex_to_rgb(stroke)
        else:
            er, eg, eb = 120, 120, 120
        conn.line.color.rgb = RGBColor(er, eg, eb)
        stroke_width = edge_style.get("strokeWidth", "1")
        conn.line.width = Pt(float(stroke_width))

        if "dashed" in edge_style or edge_style.get("dashed") == "1":
            conn.line.dash_style = 4

        # Arrowhead — use python-pptx's line end properties
        end_arrow = edge_style.get("endArrow", "")
        end_fill = edge_style.get("endFill", "0")
        if end_arrow and end_arrow != "none":
            try:
                from pptx.oxml.ns import qn as _qn
                from lxml import etree as _lx
                # Navigate to the line element within the connector
                cxnSp = conn._element
                spPr = cxnSp.find(_qn("p:spPr"))
                if spPr is None:
                    spPr = cxnSp.find(_qn("a:spPr"))
                if spPr is not None:
                    ln = spPr.find(_qn("a:ln"))
                    if ln is None:
                        ln = _lx.SubElement(spPr, _qn("a:ln"))
                    # Remove existing tailEnd if any
                    existing = ln.find(_qn("a:tailEnd"))
                    if existing is not None:
                        ln.remove(existing)
                    tail = _lx.SubElement(ln, _qn("a:tailEnd"))
                    tail.set("type", "triangle")
                    tail.set("w", "sm")
                    tail.set("len", "sm")
            except Exception:
                pass

        # Edge label — check the <object> parent for label + placeholder resolution
        edge_obj = _edge_obj  # the object wrapper if present
        edge_val = ""
        if edge_obj is not None and edge_obj.tag == "object":
            edge_val = edge_obj.get("label", "")
        if not edge_val:
            edge_val = e.get("value", "")
        if edge_val:
            lbl = strip_html(edge_val)
            # Resolve %placeholder% tokens from the object parent
            if "%" in lbl and edge_obj is not None:
                lbl = _re.sub(r"%(\w+)%", lambda m: edge_obj.get(m.group(1), m.group(1)), lbl)
            lbl = lbl.strip()
            if lbl:
                mid_x = (sx_emu + tx_emu) // 2
                mid_y = (sy_emu + ty_emu) // 2
                dx_line = abs(tx_emu - sx_emu)
                dy_line = abs(ty_emu - sy_emu)

                if dy_line > dx_line:
                    # Mostly vertical — label to the left of the line
                    # Check if midpoint overlaps any shape, and shift further if so
                    label_w = Pt(max(60, len(lbl) * 4))
                    label_h = Pt(20)
                    lbl_x = mid_x - label_w - Pt(8)
                    lbl_y = mid_y - label_h // 2
                    # Extra shift if label overlaps a shape
                    for _vid, _shp in shape_map.items():
                        if (lbl_x < _shp.left + _shp.width and lbl_x + label_w > _shp.left and
                            lbl_y < _shp.top + _shp.height and lbl_y + label_h > _shp.top):
                            lbl_x = _shp.left - label_w - Pt(8)
                            break
                else:
                    # Mostly horizontal — label centered above the connector
                    # Constrain width to the gap between endpoints
                    gap = dx_line
                    label_w = min(Pt(max(60, len(lbl) * 4)), int(gap * 0.9))
                    label_h = Pt(24)  # taller to allow 2 lines if squeezed
                    lbl_x = mid_x - label_w // 2
                    lbl_y = mid_y - label_h - Pt(2)
                txbox = slide.shapes.add_textbox(lbl_x, lbl_y, label_w, label_h)
                txbox.text_frame.word_wrap = True
                txbox.text_frame.text = lbl
                from pptx.enum.text import PP_ALIGN as _PP
                for p in txbox.text_frame.paragraphs:
                    p.alignment = _PP.CENTER
                    for r in p.runs:
                        r.font.size = Pt(7)
                        r.font.color.rgb = RGBColor(80, 80, 80)

        connector_count += 1

    prs.save(pptx_path)
    print(f"Saved {pptx_path}  ({rendered_count} shapes, {connector_count} connectors, {skipped_count} skipped)")


# ---------------------------------------------------------------------------
# drawio -> PPTX high-fidelity editable renderer
# ---------------------------------------------------------------------------

def _remove_theme_style(shp) -> None:
    element = shp._element
    style_el = element.find(qn("p:style"))
    if style_el is not None:
        element.remove(style_el)
    sp_pr = element.find(qn("p:spPr"))
    if sp_pr is not None:
        for tag in ("a:effectLst", "a:effectDag"):
            child = sp_pr.find(qn(tag))
            if child is not None:
                sp_pr.remove(child)


def _add_tail_arrow(conn) -> None:
    ln = conn._element.find(".//" + qn("a:ln"))
    if ln is None:
        return
    existing = ln.find(qn("a:tailEnd"))
    if existing is not None:
        ln.remove(existing)
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    ln.append(tail)


def _display_label(cell: ET.Element, obj: Optional[ET.Element]) -> str:
    parts: list[str] = []
    if obj is not None and obj.tag == "object":
        value = obj.get("label", "")
        if value:
            parts.append(strip_html(value))
    value = cell.get("value", "")
    if value:
        parts.append(strip_html(value))
    label = "\n".join(part for part in parts if part).strip()
    if "%" in label and obj is not None:
        label = re.sub(r"%(\w+)%", lambda m: obj.get(m.group(1), m.group(1)), label)
    return label


def _normal_hex(value: str, default: str) -> str:
    value = (value or "").strip()
    if is_hex_colour(value):
        return "#" + value.lstrip("#").upper()
    if "rgb" in value:
        nums = re.findall(r"\d+", value)
        if len(nums) >= 3:
            return rgb_to_hex(int(nums[0]), int(nums[1]), int(nums[2]))
    return default


def _style_opacity(style: dict[str, str]) -> float:
    try:
        return max(0.0, min(1.0, float(style.get("opacity", "100")) / 100.0))
    except ValueError:
        return 1.0


def _blend_with_white(rgb: tuple[int, int, int], opacity: float) -> tuple[int, int, int]:
    if opacity >= 0.999:
        return rgb
    return tuple(int(round(255 * (1.0 - opacity) + channel * opacity)) for channel in rgb)


def _is_visible(el: ET.Element, all_elements: dict[str, ET.Element]) -> bool:
    current: Optional[ET.Element] = el
    while current is not None:
        if current.get("visible") == "0":
            return False
        pid = current.get("parent", "0")
        if pid == "0" or pid not in all_elements:
            break
        current = all_elements[pid]
    return True


def _is_ellipse_style(style: dict[str, str]) -> bool:
    return "ellipse" in style or style.get("shape", "").lower() == "ellipse"


def _is_text_style(style: dict[str, str]) -> bool:
    return "text" in style and style.get("fillColor", "none").lower() in ("none", "")


def _line_end_on_box(box: tuple[int, int, int, int], other_x: int, other_y: int) -> tuple[int, int]:
    left, top, width, height = box
    cx = left + width // 2
    cy = top + height // 2
    dx_line = other_x - cx
    dy_line = other_y - cy
    if dx_line == 0 and dy_line == 0:
        return cx, cy
    half_w = width / 2
    half_h = height / 2
    sx = abs(half_w / dx_line) if dx_line else float("inf")
    sy = abs(half_h / dy_line) if dy_line else float("inf")
    scale = min(sx, sy)
    return int(cx + dx_line * scale), int(cy + dy_line * scale)


def _fidelity_font_size(
    box: tuple[int, int, int, int],
    label: str,
    is_ellipse: bool,
    is_text: bool,
    text_scale: float,
) -> float:
    _left, _top, width, height = box
    w_pt = width / 12700
    h_pt = height / 12700
    lines = [line for line in label.splitlines() if line.strip()]
    line_count = max(1, len(lines))
    longest = max((len(line) for line in label.splitlines()), default=1)
    if is_text:
        base = min(7.0, max(3.7, h_pt / (line_count * 1.45)))
    elif is_ellipse:
        base = min(5.8, max(3.2, min(w_pt / max(longest * 0.60, 4), h_pt / (line_count * 1.35))))
    else:
        base = min(7.8, max(4.0, min(w_pt / max(longest * 0.52, 7), h_pt / (line_count * 1.35))))
    return base * text_scale


def _add_fidelity_label(
    slide,
    box: tuple[int, int, int, int],
    label: str,
    style: dict[str, str],
    is_ellipse: bool,
    is_text: bool,
    text_scale: float,
) -> None:
    if not label:
        return
    left, top, width, height = box
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(1.5)
    tf.margin_bottom = Pt(1.5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if is_ellipse else MSO_ANCHOR.TOP

    align = style.get("align", "center")
    font_size = _fidelity_font_size(box, label, is_ellipse, is_text, text_scale)
    r, g, b = hex_to_rgb(_normal_hex(style.get("fontColor", ""), "#000000"))

    lines = [line.strip() for line in label.splitlines()]
    for idx, line in enumerate(lines):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT if align == "left" or is_text else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        run.font.name = "Calibri"
        run.font.size = Pt(font_size)
        run.font.bold = False
        run.font.color.rgb = RGBColor(r, g, b)


def _add_fidelity_shape(
    slide,
    box: tuple[int, int, int, int],
    style: dict[str, str],
    is_ellipse: bool,
    is_text: bool,
) -> None:
    if is_text:
        return
    left, top, width, height = box
    if is_ellipse:
        shp_type = MSO_SHAPE.OVAL
    elif style.get("rounded") == "1":
        shp_type = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        shp_type = MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shp_type, left, top, width, height)
    _remove_theme_style(shp)

    fill = style.get("fillColor", "")
    if fill.lower() == "none":
        shp.fill.background()
    else:
        fr, fg, fb = hex_to_rgb(_normal_hex(fill, "#FFFFFF"))
        fr, fg, fb = _blend_with_white((fr, fg, fb), _style_opacity(style))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(fr, fg, fb)

    stroke = style.get("strokeColor", "")
    if stroke.lower() == "none":
        shp.line.fill.background()
    else:
        lr, lg, lb = hex_to_rgb(_normal_hex(stroke, "#666666"))
        lr, lg, lb = _blend_with_white((lr, lg, lb), _style_opacity(style))
        shp.line.color.rgb = RGBColor(lr, lg, lb)
    try:
        shp.line.width = Pt(float(style.get("strokeWidth", "1") or 1) * 0.45)
    except ValueError:
        shp.line.width = Pt(0.45)
    if style.get("dashed") == "1":
        shp.line.dash_style = 4


def drawio_to_pptx_fidelity(
    drawio_path: str,
    pptx_path: str,
    exclude_layers: Optional[list[str]] = None,
    include_hidden: bool = False,
    text_scale: float = 0.82,
    margin_pt: float = 10.0,
) -> None:
    """Render visible Draw.io geometry as editable PowerPoint shapes."""
    all_elements, vertices, edges = _collect_drawio_cells(drawio_path, exclude_layers or [])
    if not include_hidden:
        vertices = [
            (cell, obj) for cell, obj in vertices
            if _is_visible(cell, all_elements) and (obj is None or _is_visible(obj, all_elements))
        ]
        edges = [
            (edge, obj) for edge, obj in edges
            if _is_visible(edge, all_elements) and (obj is None or _is_visible(obj, all_elements))
        ]

    items = []
    for cell, obj in vertices:
        geo = cell.find("mxGeometry")
        if geo is None:
            continue
        ax, ay = get_absolute_coords(all_elements, cell)
        w = float(geo.get("width", 50))
        h = float(geo.get("height", 50))
        style = parse_style(cell.get("style", ""))
        label = _display_label(cell, obj)
        vid = (obj.get("id", "") if obj is not None else "") or cell.get("id", "")
        items.append((vid, cell, obj, style, label, ax, ay, w, h))

    if not items:
        raise ValueError("No visible Draw.io vertices found")

    bb_min_x = min(item[5] for item in items)
    bb_min_y = min(item[6] for item in items)
    bb_max_x = max(item[5] + item[7] for item in items)
    bb_max_y = max(item[6] + item[8] for item in items)
    diagram_w = bb_max_x - bb_min_x
    diagram_h = bb_max_y - bb_min_y

    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    margin = Pt(margin_pt)
    scale = min(
        (prs.slide_width - 2 * margin) / (diagram_w * DRAWIO_PX_TO_EMU),
        (prs.slide_height - 2 * margin) / (diagram_h * DRAWIO_PX_TO_EMU),
    )
    offset_x = (prs.slide_width - int(diagram_w * DRAWIO_PX_TO_EMU * scale)) // 2
    offset_y = (prs.slide_height - int(diagram_h * DRAWIO_PX_TO_EMU * scale)) // 2

    def to_x(v: float) -> int:
        return int((v - bb_min_x) * DRAWIO_PX_TO_EMU * scale + offset_x)

    def to_y(v: float) -> int:
        return int((v - bb_min_y) * DRAWIO_PX_TO_EMU * scale + offset_y)

    def to_len(v: float) -> int:
        return int(v * DRAWIO_PX_TO_EMU * scale)

    boxes: dict[str, tuple[int, int, int, int]] = {}
    for vid, _cell, _obj, _style, _label, ax, ay, w, h in items:
        boxes[vid] = (to_x(ax), to_y(ay), to_len(w), to_len(h))

    for vid, _cell, _obj, style, _label, _ax, _ay, _w, _h in sorted(items, key=lambda item: item[7] * item[8], reverse=True):
        is_ellipse = _is_ellipse_style(style)
        is_text = _is_text_style(style)
        if is_ellipse:
            continue
        _add_fidelity_shape(slide, boxes[vid], style, is_ellipse, is_text)

    connector_count = 0
    for edge, obj in edges:
        geo = edge.find("mxGeometry")
        if geo is None:
            continue
        source_id = edge.get("source", "")
        target_id = edge.get("target", "")
        src_pt = geo.find("mxPoint[@as='sourcePoint']")
        tgt_pt = geo.find("mxPoint[@as='targetPoint']")

        if source_id in boxes and target_id in boxes:
            sb = boxes[source_id]
            tb = boxes[target_id]
            tcx = tb[0] + tb[2] // 2
            tcy = tb[1] + tb[3] // 2
            scx = sb[0] + sb[2] // 2
            scy = sb[1] + sb[3] // 2
            sx, sy = _line_end_on_box(sb, tcx, tcy)
            tx, ty = _line_end_on_box(tb, scx, scy)
        elif src_pt is not None and tgt_pt is not None:
            sx = to_x(float(src_pt.get("x", 0)))
            sy = to_y(float(src_pt.get("y", 0)))
            tx = to_x(float(tgt_pt.get("x", 0)))
            ty = to_y(float(tgt_pt.get("y", 0)))
        else:
            continue

        style = parse_style(edge.get("style", ""))
        conn = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, sx, sy, tx, ty)
        _remove_theme_style(conn)
        r, g, b = hex_to_rgb(_normal_hex(style.get("strokeColor", ""), "#555555"))
        conn.line.color.rgb = RGBColor(r, g, b)
        try:
            conn.line.width = Pt(float(style.get("strokeWidth", "1") or 1) * 0.45)
        except ValueError:
            conn.line.width = Pt(0.45)
        if "dashed" in style or style.get("dashed") == "1":
            conn.line.dash_style = 4
        if style.get("endArrow", "").lower() not in ("", "none"):
            _add_tail_arrow(conn)

        edge_label = _display_label(edge, obj)
        if edge_label and len(edge_label) < 70:
            mid_x = (sx + tx) // 2
            mid_y = (sy + ty) // 2
            label_w = int(max(Pt(28), min(Pt(86), len(edge_label) * Pt(3.1))))
            label_h = int(Pt(13))
            _add_fidelity_label(
                slide,
                (mid_x - label_w // 2, mid_y - label_h // 2, label_w, label_h),
                edge_label,
                {"align": "center", "fontColor": "#333333"},
                False,
                True,
                text_scale,
            )
        connector_count += 1

    for vid, _cell, _obj, style, label, _ax, _ay, _w, _h in sorted(items, key=lambda item: item[7] * item[8], reverse=True):
        is_ellipse = _is_ellipse_style(style)
        is_text = _is_text_style(style)
        if is_ellipse:
            _add_fidelity_shape(slide, boxes[vid], style, is_ellipse, is_text)
        _add_fidelity_label(slide, boxes[vid], label, style, is_ellipse, is_text, text_scale)

    prs.save(pptx_path)
    print(f"Saved {pptx_path}  ({len(items)} editable vertices, {connector_count} connectors)")


# ---------------------------------------------------------------------------
# drawio -> ODP
# ---------------------------------------------------------------------------

DRAWIO_PX_TO_CM = 0.75 * 2.54 / 72.0
SLIDE_WIDTH_CM = 25.4
SLIDE_HEIGHT_CM = 14.2875
ODP_MIMETYPE = "application/vnd.oasis.opendocument.presentation"


@dataclass
class OdpBox:
    left: float
    top: float
    width: float
    height: float


ODP_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
    "presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
    "meta": "urn:oasis:names:tc:opendocument:xmlns:meta:1.0",
    "manifest": "urn:oasis:names:tc:opendocument:xmlns:manifest:1.0",
}

for _prefix, _uri in ODP_NS.items():
    ET.register_namespace(_prefix, _uri)


def _odp_q(name: str) -> str:
    prefix, local = name.split(":", 1)
    return f"{{{ODP_NS[prefix]}}}{local}"


def _odp_attrs(attrs: Optional[dict[str, str]] = None) -> dict[str, str]:
    result: dict[str, str] = {}
    for key, value in (attrs or {}).items():
        result[_odp_q(key) if ":" in key else key] = str(value)
    return result


def _odp_el(name: str, attrs: Optional[dict[str, str]] = None) -> ET.Element:
    return ET.Element(_odp_q(name), _odp_attrs(attrs))


def _odp_child(parent: ET.Element, name: str, attrs: Optional[dict[str, str]] = None) -> ET.Element:
    return ET.SubElement(parent, _odp_q(name), _odp_attrs(attrs))


def _cm(value: float) -> str:
    return f"{value:.4f}cm"


def _pt_to_cm(value: float) -> float:
    return value * 2.54 / 72.0


def _font_colour_hex(style: dict[str, str]) -> str:
    font_colour = style.get("fontColor", "#000000")
    if is_hex_colour(font_colour):
        return "#" + font_colour.strip().lstrip("#").upper()
    if "rgb" in font_colour:
        nums = re.findall(r"\d+", font_colour)
        if len(nums) >= 3:
            return rgb_to_hex(int(nums[0]), int(nums[1]), int(nums[2]))
    return "#000000"


def _collect_drawio_cells(
    drawio_path: str,
    exclude_layers: Optional[list[str]] = None,
) -> tuple[dict[str, ET.Element], list[tuple[ET.Element, Optional[ET.Element]]], list[tuple[ET.Element, Optional[ET.Element]]]]:
    exclude_layers = [l.lower() for l in (exclude_layers or [])]

    mg = load_drawio_xml(drawio_path)
    root_el = mg.find("root")
    if root_el is None:
        raise ValueError("No <root> element in mxGraphModel")

    all_elements: dict[str, ET.Element] = {}
    for el in root_el.iter():
        eid = el.get("id")
        if eid:
            all_elements[eid] = el

    def is_excluded(el: ET.Element) -> bool:
        current = el
        while current is not None:
            val = (current.get("value") or "").strip().lower()
            if val and val in exclude_layers:
                return True
            pid = current.get("parent", "0")
            if pid == "0" or pid not in all_elements:
                break
            current = all_elements[pid]
        return False

    vertices: list[tuple[ET.Element, Optional[ET.Element]]] = []
    edges: list[tuple[ET.Element, Optional[ET.Element]]] = []
    seen_cells: set[int] = set()

    for obj in root_el.iter("object"):
        if is_excluded(obj):
            continue
        for cell in obj.findall("mxCell"):
            if id(cell) in seen_cells:
                continue
            seen_cells.add(id(cell))
            if cell.get("vertex") == "1":
                vertices.append((cell, obj))
            elif cell.get("edge") == "1":
                edges.append((cell, obj))

    for el in root_el.iter("mxCell"):
        if id(el) in seen_cells:
            continue
        seen_cells.add(id(el))
        if is_excluded(el):
            continue
        if el.get("vertex") == "1":
            vertices.append((el, None))
        elif el.get("edge") == "1":
            edges.append((el, None))

    return all_elements, vertices, edges


def _add_graphic_style(
    auto_styles: ET.Element,
    name: str,
    fill: Optional[str],
    stroke: Optional[str],
    stroke_width_cm: float = 0.0265,
    dashed: bool = False,
    marker_end: bool = False,
    text_align: str = "center",
    vertical_align: str = "middle",
) -> None:
    style_el = _odp_child(
        auto_styles,
        "style:style",
        {"style:name": name, "style:family": "graphic"},
    )
    props = {
        "draw:fill": "none" if fill is None else "solid",
        "draw:stroke": "none" if stroke is None else ("dash" if dashed else "solid"),
        "draw:textarea-horizontal-align": text_align,
        "draw:textarea-vertical-align": vertical_align,
        "fo:padding": "0.07cm",
    }
    if fill is not None:
        props["draw:fill-color"] = fill
    if stroke is not None:
        props["svg:stroke-color"] = stroke
        props["svg:stroke-width"] = _cm(stroke_width_cm)
    if dashed:
        props["draw:stroke-dash"] = "dash1"
    if marker_end:
        props["draw:marker-end"] = "Triangle"
        props["draw:marker-end-width"] = "0.18cm"
        props["draw:marker-end-center"] = "false"
    _odp_child(style_el, "style:graphic-properties", props)


def _add_paragraph_style(auto_styles: ET.Element, name: str, align: str) -> None:
    style_el = _odp_child(
        auto_styles,
        "style:style",
        {"style:name": name, "style:family": "paragraph"},
    )
    _odp_child(style_el, "style:paragraph-properties", {"fo:text-align": align})


def _add_text_style(
    auto_styles: ET.Element,
    name: str,
    size_pt: float,
    colour: str,
    bold: bool = False,
) -> None:
    style_el = _odp_child(
        auto_styles,
        "style:style",
        {"style:name": name, "style:family": "text"},
    )
    props = {
        "fo:font-size": f"{max(size_pt, 1):.1f}pt",
        "fo:color": colour,
    }
    if bold:
        props["fo:font-weight"] = "bold"
    _odp_child(style_el, "style:text-properties", props)


def _append_text_lines(
    parent: ET.Element,
    auto_styles: ET.Element,
    label: str,
    style: dict[str, str],
    scale: float,
    style_counter: list[int],
    into_text_box: bool = False,
) -> None:
    container = parent
    if into_text_box:
        container = _odp_child(parent, "draw:text-box")

    lines = [line.strip() for line in label.splitlines() if line.strip()]
    if not lines:
        _odp_child(container, "text:p", {"text:style-name": "Pcenter"})
        return

    font_colour = _font_colour_hex(style)
    raw_html = style.get("_raw_html", "")
    font_sizes = re.findall(r"font-size:\s*(\d+)px", raw_html)
    align = "Pleft" if style.get("align") == "left" else "Pcenter"

    for idx, line in enumerate(lines):
        p_el = _odp_child(container, "text:p", {"text:style-name": align})
        style_counter[0] += 1
        text_style = f"T{style_counter[0]}"
        if idx == 0:
            size_pt = (float(font_sizes[0]) if font_sizes else 16.0) * scale * 0.75
            _add_text_style(auto_styles, text_style, size_pt, font_colour, bold=True)
        elif line.startswith("["):
            _add_text_style(auto_styles, text_style, 8.0, font_colour)
        else:
            size_pt = (float(font_sizes[-1]) if len(font_sizes) > 1 else 11.0) * scale * 0.75
            if font_colour.upper() == "#FFFFFF":
                desc_colour = "#CCCCCC"
            else:
                r, g, b = hex_to_rgb(font_colour)
                desc_colour = rgb_to_hex(min(r + 60, 255), min(g + 60, 255), min(b + 60, 255))
            _add_text_style(auto_styles, text_style, size_pt, desc_colour)
        span = _odp_child(p_el, "text:span", {"text:style-name": text_style})
        span.text = line


def _xml_document_bytes(root: ET.Element) -> bytes:
    ET.indent(root, space="  ")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _build_odp_styles_xml() -> bytes:
    root = _odp_el("office:document-styles", {"office:version": "1.2"})
    _odp_child(root, "office:font-face-decls")
    _odp_child(root, "office:styles")
    auto_styles = _odp_child(root, "office:automatic-styles")
    page_layout = _odp_child(auto_styles, "style:page-layout", {"style:name": "PM1"})
    _odp_child(
        page_layout,
        "style:page-layout-properties",
        {
            "fo:page-width": _cm(SLIDE_WIDTH_CM),
            "fo:page-height": _cm(SLIDE_HEIGHT_CM),
            "style:print-orientation": "landscape",
        },
    )
    master_styles = _odp_child(root, "office:master-styles")
    _odp_child(
        master_styles,
        "style:master-page",
        {"style:name": "Default", "style:page-layout-name": "PM1"},
    )
    return _xml_document_bytes(root)


def _build_odp_meta_xml() -> bytes:
    root = _odp_el("office:document-meta", {"office:version": "1.2"})
    meta_el = _odp_child(root, "office:meta")
    generator = _odp_child(meta_el, "meta:generator")
    generator.text = "drawio_ppt.py"
    created = _odp_child(meta_el, "meta:creation-date")
    created.text = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    return _xml_document_bytes(root)


def _build_odp_settings_xml() -> bytes:
    root = _odp_el("office:document-settings", {"office:version": "1.2"})
    _odp_child(root, "office:settings")
    return _xml_document_bytes(root)


def _build_odp_manifest_xml() -> bytes:
    root = _odp_el("manifest:manifest", {"manifest:version": "1.2"})
    entries = [
        ("/", ODP_MIMETYPE),
        ("content.xml", "text/xml"),
        ("styles.xml", "text/xml"),
        ("meta.xml", "text/xml"),
        ("settings.xml", "text/xml"),
    ]
    for path, media_type in entries:
        _odp_child(
            root,
            "manifest:file-entry",
            {"manifest:full-path": path, "manifest:media-type": media_type},
        )
    return _xml_document_bytes(root)


def drawio_to_odp(
    drawio_path: str,
    odp_path: str,
    config_path: Optional[str] = None,
    exclude_layers: Optional[list[str]] = None,
    back_layers: Optional[list[str]] = None,
    line_darken: int = 20,
) -> None:
    back_layers = [l.lower() for l in (back_layers or [])]
    palette = load_palette(config_path) if config_path else []
    all_elements, vertices, edges = _collect_drawio_cells(drawio_path, exclude_layers)

    def is_back_layer(el: ET.Element) -> bool:
        current = el
        while current is not None:
            val = (current.get("value") or "").strip().lower()
            if val and val in back_layers:
                return True
            pid = current.get("parent", "0")
            if pid == "0" or pid not in all_elements:
                break
            current = all_elements[pid]
        return False

    bb_min_x, bb_min_y = float("inf"), float("inf")
    bb_max_x, bb_max_y = float("-inf"), float("-inf")
    for v, _obj in vertices:
        geo = v.find("mxGeometry")
        if geo is None:
            continue
        ax, ay = get_absolute_coords(all_elements, v)
        w = float(geo.get("width", 50))
        h = float(geo.get("height", 50))
        bb_min_x = min(bb_min_x, ax)
        bb_min_y = min(bb_min_y, ay)
        bb_max_x = max(bb_max_x, ax + w)
        bb_max_y = max(bb_max_y, ay + h)

    if not math.isfinite(bb_min_x) or not math.isfinite(bb_min_y):
        raise ValueError("No renderable vertices found in Draw.io file")

    diag_w = bb_max_x - bb_min_x or 1
    diag_h = bb_max_y - bb_min_y or 1
    margin_cm = _pt_to_cm(30)
    avail_w = SLIDE_WIDTH_CM - 2 * margin_cm
    avail_h = SLIDE_HEIGHT_CM - 2 * margin_cm
    scale = min(
        avail_w / (diag_w * DRAWIO_PX_TO_CM),
        avail_h / (diag_h * DRAWIO_PX_TO_CM),
        1.0,
    )
    rendered_w = diag_w * DRAWIO_PX_TO_CM * scale
    rendered_h = diag_h * DRAWIO_PX_TO_CM * scale
    offset_x = (SLIDE_WIDTH_CM - rendered_w) / 2
    offset_y = (SLIDE_HEIGHT_CM - rendered_h) / 2

    def to_cm_x(val: float) -> float:
        return (val - bb_min_x) * DRAWIO_PX_TO_CM * scale + offset_x

    def to_cm_y(val: float) -> float:
        return (val - bb_min_y) * DRAWIO_PX_TO_CM * scale + offset_y

    def to_cm_len(val: float) -> float:
        return val * DRAWIO_PX_TO_CM * scale

    content = _odp_el("office:document-content", {"office:version": "1.2"})
    _odp_child(content, "office:scripts")
    _odp_child(content, "office:font-face-decls")
    auto_styles = _odp_child(content, "office:automatic-styles")
    _odp_child(
        auto_styles,
        "draw:marker",
        {
            "draw:name": "Triangle",
            "svg:viewBox": "0 0 20 30",
            "svg:d": "M10 0 L20 30 L0 30 Z",
        },
    )
    _odp_child(
        auto_styles,
        "draw:stroke-dash",
        {
            "draw:name": "dash1",
            "draw:style": "rect",
            "draw:dots1": "1",
            "draw:dots1-length": "0.12cm",
            "draw:distance": "0.08cm",
        },
    )
    page_style = _odp_child(
        auto_styles,
        "style:style",
        {"style:name": "dp1", "style:family": "drawing-page"},
    )
    _odp_child(
        page_style,
        "style:drawing-page-properties",
        {"draw:fill": "solid", "draw:fill-color": "#FFFFFF"},
    )
    _add_paragraph_style(auto_styles, "Pcenter", "center")
    _add_paragraph_style(auto_styles, "Pleft", "left")

    body = _odp_child(content, "office:body")
    presentation = _odp_child(body, "office:presentation")
    page = _odp_child(
        presentation,
        "draw:page",
        {"draw:name": "page1", "draw:style-name": "dp1", "draw:master-page-name": "Default"},
    )

    def _shape_area(item: tuple[ET.Element, Optional[ET.Element]]) -> float:
        v, _ = item
        geo = v.find("mxGeometry")
        if geo is None:
            return 0.0
        return float(geo.get("width", 50)) * float(geo.get("height", 50))

    vertices.sort(key=lambda item: (0 if is_back_layer(item[0]) else 1, -_shape_area(item)))
    shape_map: dict[str, OdpBox] = {}
    rendered_count = 0
    skipped_count = 0
    connector_count = 0
    style_counter = [0]
    graphic_style_counter = 0
    z_index = 0
    min_shape_area = 1500

    for v, obj_parent in vertices:
        geo = v.find("mxGeometry")
        if geo is None:
            continue
        style = parse_style(v.get("style", ""))
        raw_html = (obj_parent.get("label", "") if obj_parent is not None else "") or v.get("value", "") or ""
        style["_raw_html"] = raw_html
        label = resolve_label(v, obj_parent)
        ax, ay = get_absolute_coords(all_elements, v)
        w = float(geo.get("width", 50))
        h = float(geo.get("height", 50))

        if "ellipse" in style:
            skipped_count += 1
            continue
        if w * h < min_shape_area and not label.strip():
            skipped_count += 1
            continue

        vid = (obj_parent.get("id", "") if obj_parent is not None else "") or v.get("id", "")
        box = OdpBox(to_cm_x(ax), to_cm_y(ay), to_cm_len(w), to_cm_len(h))
        is_text_only = "text" in style and style.get("fillColor", "none").lower() in ("none", "")

        raw_fill = style.get("fillColor", "")
        if is_text_only or raw_fill.lower() == "none":
            fill_hex = None
        elif is_hex_colour(raw_fill):
            fill_hex = "#" + raw_fill.strip().lstrip("#").upper()
            if palette:
                fill_hex = snap_colour(fill_hex, palette)
        else:
            fill_hex = "#FFFFFF"

        raw_stroke = style.get("strokeColor", "")
        if is_text_only or raw_stroke.lower() == "none":
            stroke_hex = None
        elif is_hex_colour(raw_stroke):
            stroke_hex = "#" + raw_stroke.strip().lstrip("#").upper()
            if line_darken > 0 and palette:
                stroke_hex = darken(stroke_hex, line_darken)
        else:
            stroke_hex = "#000000"

        graphic_style_counter += 1
        graphic_style = f"gr{graphic_style_counter}"
        _add_graphic_style(
            auto_styles,
            graphic_style,
            fill_hex,
            stroke_hex,
            dashed=style.get("dashed") == "1",
            text_align="left" if style.get("align") == "left" else "center",
            vertical_align="top" if style.get("verticalAlign") == "top" else "middle",
        )

        z_index += 1
        attrs = {
            "draw:name": f"shape{rendered_count + 1}",
            "draw:style-name": graphic_style,
            "draw:z-index": str(z_index),
            "svg:x": _cm(box.left),
            "svg:y": _cm(box.top),
            "svg:width": _cm(box.width),
            "svg:height": _cm(box.height),
        }
        if is_text_only:
            shape_el = _odp_child(page, "draw:frame", attrs)
            _append_text_lines(shape_el, auto_styles, label, style, scale, style_counter, into_text_box=True)
        else:
            if style.get("rounded") == "1":
                attrs["draw:corner-radius"] = _cm(min(box.width, box.height) * 0.08)
            shape_el = _odp_child(page, "draw:rect", attrs)
            _append_text_lines(shape_el, auto_styles, label, style, scale, style_counter)

        shape_map[vid] = box
        rendered_count += 1

    simple_diagram = rendered_count <= 30

    def _edge_intersect(box: OdpBox, other_cx: float, other_cy: float) -> tuple[float, float]:
        cx = box.left + box.width / 2
        cy = box.top + box.height / 2
        ddx = other_cx - cx
        ddy = other_cy - cy
        if ddx == 0 and ddy == 0:
            return cx, cy
        hw = box.width / 2
        hh = box.height / 2
        s_x = abs(hw / ddx) if ddx != 0 else float("inf")
        s_y = abs(hh / ddy) if ddy != 0 else float("inf")
        s = min(s_x, s_y)
        return cx + ddx * s, cy + ddy * s

    def _best_edge_pair(source: OdpBox, target: OdpBox) -> tuple[tuple[float, float], tuple[float, float]]:
        sc = (source.left + source.width / 2, source.top + source.height / 2)
        tc = (target.left + target.width / 2, target.top + target.height / 2)
        dx_abs = abs(tc[0] - sc[0])
        dy_abs = abs(tc[1] - sc[1])
        if dy_abs >= dx_abs * 0.5:
            src_x = min(max(tc[0], source.left), source.left + source.width)
            if sc[1] < tc[1]:
                return (src_x, source.top + source.height), (tc[0], target.top)
            return (src_x, source.top), (tc[0], target.top + target.height)
        if sc[0] < tc[0]:
            return (source.left + source.width, sc[1]), (target.left, tc[1])
        return (source.left, sc[1]), (target.left + target.width, tc[1])

    for e, edge_obj in edges:
        geo = e.find("mxGeometry")
        if geo is None:
            continue

        source_id = e.get("source", "")
        target_id = e.get("target", "")
        src_pt = geo.find("mxPoint[@as='sourcePoint']")
        tgt_pt = geo.find("mxPoint[@as='targetPoint']")

        if src_pt is not None:
            sx = to_cm_x(float(src_pt.get("x", 0)))
            sy = to_cm_y(float(src_pt.get("y", 0)))
            if tgt_pt is not None:
                tx = to_cm_x(float(tgt_pt.get("x", 0)))
                ty = to_cm_y(float(tgt_pt.get("y", 0)))
            elif simple_diagram and target_id in shape_map:
                tx, ty = _edge_intersect(shape_map[target_id], sx, sy)
            else:
                continue
        elif simple_diagram and source_id in shape_map:
            source = shape_map[source_id]
            if target_id in shape_map:
                (sx, sy), (tx, ty) = _best_edge_pair(source, shape_map[target_id])
            elif tgt_pt is not None:
                tx = to_cm_x(float(tgt_pt.get("x", 0)))
                ty = to_cm_y(float(tgt_pt.get("y", 0)))
                sx, sy = _edge_intersect(source, tx, ty)
            else:
                continue
        else:
            continue

        edge_style = parse_style(e.get("style", ""))
        stroke = edge_style.get("strokeColor", "")
        stroke_hex = "#" + stroke.strip().lstrip("#").upper() if is_hex_colour(stroke) else "#787878"
        stroke_width = _pt_to_cm(float(edge_style.get("strokeWidth", "1") or 1))
        graphic_style_counter += 1
        line_style = f"gr{graphic_style_counter}"
        _add_graphic_style(
            auto_styles,
            line_style,
            None,
            stroke_hex,
            stroke_width_cm=stroke_width,
            dashed=("dashed" in edge_style or edge_style.get("dashed") == "1"),
            marker_end=edge_style.get("endArrow", "") not in ("", "none"),
        )
        z_index += 1
        _odp_child(
            page,
            "draw:line",
            {
                "draw:name": f"connector{connector_count + 1}",
                "draw:style-name": line_style,
                "draw:z-index": str(z_index),
                "svg:x1": _cm(sx),
                "svg:y1": _cm(sy),
                "svg:x2": _cm(tx),
                "svg:y2": _cm(ty),
            },
        )

        edge_val = edge_obj.get("label", "") if edge_obj is not None and edge_obj.tag == "object" else ""
        if not edge_val:
            edge_val = e.get("value", "")
        if edge_val:
            lbl = strip_html(edge_val)
            if "%" in lbl and edge_obj is not None:
                lbl = re.sub(r"%(\w+)%", lambda m: edge_obj.get(m.group(1), m.group(1)), lbl)
            lbl = lbl.strip()
            if lbl:
                dx_line = abs(tx - sx)
                dy_line = abs(ty - sy)
                label_w = min(max(1.6, len(lbl) * 0.11), max(1.6, dx_line * 0.9 if dx_line else 2.5))
                label_h = 0.65
                if dy_line > dx_line:
                    label_x = (sx + tx) / 2 - label_w - 0.2
                    label_y = (sy + ty) / 2 - label_h / 2
                else:
                    label_x = (sx + tx) / 2 - label_w / 2
                    label_y = (sy + ty) / 2 - label_h - 0.05
                graphic_style_counter += 1
                label_style = f"gr{graphic_style_counter}"
                _add_graphic_style(auto_styles, label_style, None, None)
                z_index += 1
                frame = _odp_child(
                    page,
                    "draw:frame",
                    {
                        "draw:name": f"connectorLabel{connector_count + 1}",
                        "draw:style-name": label_style,
                        "draw:z-index": str(z_index),
                        "svg:x": _cm(label_x),
                        "svg:y": _cm(label_y),
                        "svg:width": _cm(label_w),
                        "svg:height": _cm(label_h),
                    },
                )
                _add_text_style(auto_styles, f"Tedge{connector_count + 1}", 7.0, "#505050")
                text_box = _odp_child(frame, "draw:text-box")
                p_el = _odp_child(text_box, "text:p", {"text:style-name": "Pcenter"})
                span = _odp_child(p_el, "text:span", {"text:style-name": f"Tedge{connector_count + 1}"})
                span.text = lbl

        connector_count += 1

    content_xml = _xml_document_bytes(content)

    odp_file = Path(odp_path)
    odp_file.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(odp_file, "w") as zf:
        info = zipfile.ZipInfo("mimetype")
        info.compress_type = zipfile.ZIP_STORED
        zf.writestr(info, ODP_MIMETYPE.encode("ascii"))
        zf.writestr("content.xml", content_xml, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("styles.xml", _build_odp_styles_xml(), compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("meta.xml", _build_odp_meta_xml(), compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("settings.xml", _build_odp_settings_xml(), compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("META-INF/manifest.xml", _build_odp_manifest_xml(), compress_type=zipfile.ZIP_DEFLATED)

    print(f"Saved {odp_path}  ({rendered_count} shapes, {connector_count} connectors, {skipped_count} skipped)")


# ---------------------------------------------------------------------------
# pptx → drawio
# ---------------------------------------------------------------------------

EMU_TO_DRAWIO_PX = 1.0 / DRAWIO_PX_TO_EMU  # inverse


def _emu_to_dx(val: int) -> float:
    return round(val * EMU_TO_DRAWIO_PX, 1)


def _pptx_color_hex(color_format) -> str:
    try:
        rgb = color_format.rgb
        if rgb:
            return f"#{rgb}"
    except Exception:
        pass
    return ""


def _shape_to_style(shp) -> str:
    props: dict[str, str] = {}

    # Shape type
    try:
        auto_type = shp.auto_shape_type
        if auto_type == MSO_SHAPE.OVAL:
            props["ellipse"] = ""
        elif auto_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            props["rounded"] = "1"
    except Exception:
        pass

    props["whiteSpace"] = "wrap"
    props["html"] = "1"

    # Fill
    try:
        if shp.fill.type is not None:
            hex_c = _pptx_color_hex(shp.fill.fore_color)
            if hex_c:
                props["fillColor"] = hex_c
    except Exception:
        props["fillColor"] = "none"

    # Stroke
    try:
        hex_c = _pptx_color_hex(shp.line.color)
        if hex_c:
            props["strokeColor"] = hex_c
    except Exception:
        pass

    return build_style(props)


def pptx_to_drawio(pptx_path: str, drawio_path: str) -> None:
    prs = Presentation(pptx_path)

    mxfile = ET.Element("mxfile")
    mxfile.set("host", "drawio_ppt.py")

    for slide_idx, slide in enumerate(prs.slides):
        diagram = ET.SubElement(mxfile, "diagram")
        diagram.set("name", f"Slide {slide_idx + 1}")
        diagram.set("id", f"slide_{slide_idx}")

        mg = ET.SubElement(diagram, "mxGraphModel")
        mg.set("dx", "0")
        mg.set("dy", "0")
        mg.set("grid", "1")
        mg.set("gridSize", "10")
        mg.set("guides", "1")
        mg.set("tooltips", "1")
        mg.set("connect", "1")
        mg.set("arrows", "1")
        mg.set("fold", "1")
        mg.set("page", "1")
        mg.set("pageScale", "1")
        mg.set("pageWidth", str(int(_emu_to_dx(prs.slide_width))))
        mg.set("pageHeight", str(int(_emu_to_dx(prs.slide_height))))

        root = ET.SubElement(mg, "root")
        ET.SubElement(root, "mxCell", id="0")
        ET.SubElement(root, "mxCell", id="1", parent="0")

        shape_id_map: dict[int, str] = {}  # pptx shape.shape_id → drawio id
        next_id = 2

        for shp in slide.shapes:
            did = str(next_id)
            next_id += 1
            shape_id_map[shp.shape_id] = did

            if shp.shape_type == MSO_SHAPE_TYPE.FREEFORM or shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue  # skip complex shapes

            # Check if it's a connector
            if hasattr(shp, "begin_x") and hasattr(shp, "end_x"):
                cell = ET.SubElement(root, "mxCell")
                cell.set("id", did)
                cell.set("value", "")
                cell.set("edge", "1")
                cell.set("parent", "1")
                cell.set("style", "endArrow=classic;html=1;")

                geo = ET.SubElement(cell, "mxGeometry")
                geo.set("relative", "1")
                geo.set("as", "geometry")

                src = ET.SubElement(geo, "mxPoint")
                src.set("x", str(_emu_to_dx(shp.begin_x)))
                src.set("y", str(_emu_to_dx(shp.begin_y)))
                src.set("as", "sourcePoint")

                tgt = ET.SubElement(geo, "mxPoint")
                tgt.set("x", str(_emu_to_dx(shp.end_x)))
                tgt.set("y", str(_emu_to_dx(shp.end_y)))
                tgt.set("as", "targetPoint")
                continue

            # Regular shape (vertex)
            label = ""
            if shp.has_text_frame:
                label = shp.text_frame.text

            cell = ET.SubElement(root, "mxCell")
            cell.set("id", did)
            cell.set("value", label)
            cell.set("style", _shape_to_style(shp))
            cell.set("vertex", "1")
            cell.set("parent", "1")

            geo = ET.SubElement(cell, "mxGeometry")
            geo.set("x", str(_emu_to_dx(shp.left)))
            geo.set("y", str(_emu_to_dx(shp.top)))
            geo.set("width", str(_emu_to_dx(shp.width)))
            geo.set("height", str(_emu_to_dx(shp.height)))
            geo.set("as", "geometry")

    tree = ET.ElementTree(mxfile)
    ET.indent(tree, space="  ")
    tree.write(drawio_path, encoding="unicode", xml_declaration=True)
    total_shapes = sum(len(s.shapes) for s in prs.slides)
    print(f"Saved {drawio_path}  ({len(prs.slides)} slides, {total_shapes} shapes)")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Draw.io <-> PowerPoint/OpenDocument converter",
    )
    sub = parser.add_subparsers(dest="command")

    d2p = sub.add_parser("drawio2ppt", help="Convert .drawio XML to .pptx")
    d2p.add_argument("input", help="Path to .drawio or .drawio.xml file")
    d2p.add_argument("output", help="Output .pptx path")
    d2p.add_argument("--config", help="Path to config.ini with [colours] palette")
    d2p.add_argument("--exclude-layers", nargs="*", default=[], help="Layer names to exclude")
    d2p.add_argument("--back-layers", nargs="*", default=[], help="Layer names to send to back")
    d2p.add_argument("--line-darken", type=int, default=20, help="Darken stroke by N (default 20)")

    d2pf = sub.add_parser("drawio2ppt-fidelity", help="Convert .drawio XML to an editable high-fidelity .pptx")
    d2pf.add_argument("input", help="Path to .drawio or .drawio.xml file")
    d2pf.add_argument("output", help="Output .pptx path")
    d2pf.add_argument("--exclude-layers", nargs="*", default=[], help="Layer names to exclude")
    d2pf.add_argument("--include-hidden", action="store_true", help="Render cells under hidden Draw.io layers")
    d2pf.add_argument("--text-scale", type=float, default=0.82, help="Multiplier for generated text size")
    d2pf.add_argument("--margin-pt", type=float, default=10.0, help="Slide margin in points")

    d2o = sub.add_parser("drawio2odp", help="Convert .drawio XML to .odp without Office automation")
    d2o.add_argument("input", help="Path to .drawio or .drawio.xml file")
    d2o.add_argument("output", help="Output .odp path")
    d2o.add_argument("--config", help="Path to config.ini with [colours] palette")
    d2o.add_argument("--exclude-layers", nargs="*", default=[], help="Layer names to exclude")
    d2o.add_argument("--back-layers", nargs="*", default=[], help="Layer names to send to back")
    d2o.add_argument("--line-darken", type=int, default=20, help="Darken stroke by N (default 20)")

    p2d = sub.add_parser("ppt2drawio", help="Convert .pptx to .drawio")
    p2d.add_argument("input", help="Path to .pptx file")
    p2d.add_argument("output", help="Output .drawio path")

    args = parser.parse_args()

    if args.command == "drawio2ppt":
        drawio_to_pptx(
            args.input, args.output,
            config_path=args.config,
            exclude_layers=args.exclude_layers,
            back_layers=args.back_layers,
            line_darken=args.line_darken,
        )
    elif args.command == "drawio2ppt-fidelity":
        drawio_to_pptx_fidelity(
            args.input,
            args.output,
            exclude_layers=args.exclude_layers,
            include_hidden=args.include_hidden,
            text_scale=args.text_scale,
            margin_pt=args.margin_pt,
        )
    elif args.command == "drawio2odp":
        drawio_to_odp(
            args.input, args.output,
            config_path=args.config,
            exclude_layers=args.exclude_layers,
            back_layers=args.back_layers,
            line_darken=args.line_darken,
        )
    elif args.command == "ppt2drawio":
        pptx_to_drawio(args.input, args.output)
    else:
        parser.print_help()
        sys.exit(1)


if __name__ == "__main__":
    main()
